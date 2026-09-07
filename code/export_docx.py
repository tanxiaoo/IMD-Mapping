"""Export report/report.md to .docx.

Written for the review export in CLAUDE.md. The one thing it must get right is
image resolution: the report embeds figures with paths relative to
`report/report.md` itself -- `figs/x.png` for the built figures and
`../outputs_*/x.png` for the notebook ones -- so every path is resolved against
the MARKDOWN FILE'S directory, not the working directory. Resolving against the
cwd silently produces a document with the built figures present and the notebook
figures missing, which is the failure this script exists to make impossible: a
path that does not resolve raises rather than being skipped.

Supports the subset of markdown the report actually uses: ATX headings, GFM
pipe tables, images with `![Figure N](path)` alt text, paragraphs, `-` bullet
lists, and inline `**bold**` / `*italic*` / `` `code` ``.

Usage:  python code/export_docx.py                       # whole report
        python code/export_docx.py --sections 4 6        # only those top-level
                                                         # sections
        python code/export_docx.py --out path/to/x.docx
"""

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_REPORT = os.path.join(REPO, 'report', 'report.md')

# Widest a figure may be laid out at. The report's figures vary from 8 to 15
# inches of source width, so they are scaled to fit the page rather than
# inserted at native size.
MAX_FIG_IN = 6.2


def load_sections(path, wanted):
    """Report lines, optionally restricted to some top-level sections.

    `wanted` is a set of section numbers as strings ('4', '6'). The title and
    everything before the first `## N.` heading are kept only for a full
    export; a section subset starts at its own heading, since a partial
    document is for checking rendering, not for reading as a report.
    """
    lines = open(path, encoding='utf-8').read().split('\n')
    if not wanted:
        return lines

    out, keep = [], False
    for line in lines:
        m = re.match(r'^## (\d+)\.', line)
        if m:
            keep = m.group(1) in wanted
        if keep:
            out.append(line)
    return out


def add_runs(par, text):
    """Inline **bold**, *italic* and `code` into a paragraph."""
    for part in re.split(r'(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)',
                         text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
        elif part.startswith('*') and part.endswith('*'):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _row_cells(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def _is_separator(line):
    cells = _row_cells(line)
    return bool(cells) and all(set(c) <= set('-: ') and c for c in cells)


def add_table(doc, rows):
    """A GFM pipe table as a Word table with a bold header row."""
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell in enumerate(header):
        par = t.rows[0].cells[j].paragraphs[0]
        add_runs(par, cell)
        for run in par.runs:
            run.bold = True
            run.font.size = Pt(9)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            par = t.rows[i].cells[j].paragraphs[0]
            add_runs(par, row[j] if j < len(row) else '')
            for run in par.runs:
                run.font.size = Pt(9)
    return t


def add_image(doc, md_dir, path, report):
    """Insert one figure, resolving `path` against the markdown's directory.

    Raises on an unresolvable path. Skipping a missing figure would produce a
    document that looks complete and is not, which is exactly the failure a
    relative-path export is prone to.
    """
    resolved = os.path.normpath(os.path.join(md_dir, path))
    if not os.path.exists(resolved):
        raise SystemExit(
            f'Figure not found: {path!r} resolved to {resolved!r}. Image '
            f'paths are relative to the markdown file ({md_dir}), not to the '
            f'working directory.')

    from PIL import Image
    with Image.open(resolved) as im:
        w_px, h_px = im.size
        dpi = im.info.get('dpi', (150, 150))[0] or 150
    width_in = min(MAX_FIG_IN, w_px / dpi)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(resolved, width=Inches(width_in))

    report.append(dict(path=path, resolved=_rel(resolved),
                       px=f'{w_px}x{h_px}', dpi=int(dpi),
                       placed_in=round(width_in, 2)))


def convert(lines, md_dir, out_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    figures = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        m = re.match(r'^(#{1,6})\s+(.*)', stripped)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            i += 1
            continue

        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if m:
            add_image(doc, md_dir, m.group(2), figures)
            i += 1
            continue

        # Pipe table: a header row followed by a separator row.
        if stripped.startswith('|') and i + 1 < n and _is_separator(lines[i + 1]):
            rows = [_row_cells(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(_row_cells(lines[i]))
                i += 1
            add_table(doc, rows)
            doc.add_paragraph()
            continue

        if re.match(r'^[-*]\s+', stripped):
            while i < n and re.match(r'^[-*]\s+', lines[i].strip()):
                par = doc.add_paragraph(style='List Bullet')
                add_runs(par, re.sub(r'^[-*]\s+', '', lines[i].strip()))
                i += 1
            continue

        # Paragraph: consume until a blank line or a block-level construct.
        buf = []
        while i < n and lines[i].strip():
            s = lines[i].strip()
            if re.match(r'^(#{1,6}\s|!\[|\|)', s) or re.match(r'^[-*]\s+', s):
                break
            buf.append(s)
            i += 1
        if buf:
            text = ' '.join(buf)
            par = doc.add_paragraph()
            add_runs(par, text)
            # Figure captions read as captions, not as body text.
            if re.match(r'^\*{0,2}(Figure|Table)\s+\d+[.·]', text):
                for run in par.runs:
                    run.font.size = Pt(9)
                    run.italic = True

    doc.save(out_path)
    return figures


def _rel(path):
    """Path relative to the repo, or as given if it is on another drive.

    os.path.relpath raises across Windows drive letters, which happens
    whenever the output is written to a scratch directory on another volume.
    """
    try:
        return os.path.relpath(path, REPO)
    except ValueError:
        return path


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--report', default=DEFAULT_REPORT)
    ap.add_argument('--out', default=None)
    ap.add_argument('--sections', nargs='*', default=None,
                    help='top-level section numbers, e.g. --sections 4 6')
    args = ap.parse_args()

    md_dir = os.path.dirname(os.path.abspath(args.report))
    wanted = set(args.sections or [])
    lines = load_sections(args.report, wanted)

    if not lines:
        raise SystemExit(f'No content selected from {args.report}'
                         + (f' for sections {sorted(wanted)}' if wanted else ''))

    out = args.out or os.path.join(
        md_dir, 'report' + (f'_s{"".join(sorted(wanted))}' if wanted else '')
        + '.docx')

    figures = convert(lines, md_dir, out)

    width = 74
    print('=' * width)
    print('  export_docx.py')
    print('=' * width)
    print(f'source  : {_rel(args.report)}')
    print(f'sections: {" ".join(sorted(wanted)) if wanted else "all"}')
    print(f'lines   : {len(lines):,}')
    print(f'output  : {_rel(out)}')
    print()
    print(f'Figures embedded: {len(figures)}  '
          f'(image paths resolved against {_rel(md_dir)}/)')
    print()
    if figures:
        print(f'  {"as written in markdown":<46} {"px":>11} {"dpi":>4} '
              f'{"placed":>7}')
        print('  ' + '-' * (width - 4))
        for f in figures:
            print(f'  {f["path"][:46]:<46} {f["px"]:>11} {f["dpi"]:>4} '
                  f'{f["placed_in"]:>6}"')
        print()
        outside = [f for f in figures if f['path'].startswith('..')]
        print(f'  of which use ../ to leave report/ : {len(outside)}')
        for f in outside:
            print(f'    {f["path"]}  ->  {f["resolved"]}')
    print()
    print(f'Wrote {os.path.getsize(out):,} bytes')
    print('=' * width)
    return 0


if __name__ == '__main__':
    sys.exit(main())
